"""
Ground Truth vs SegResNet evaluation for the BraTS 2023 test set.

Run from:
C:/Users/chand/BrainTumorProject/Brain-Tumors-Segmentation

    conda activate brats_segmentation
    python evaluate_ground_truth_vs_prediction.py

Creates:
    evaluation_results/ground_truth_vs_prediction.csv
    evaluation_results/ground_truth_vs_prediction.xlsx
"""
import csv
import gc
import sys
from pathlib import Path

import numpy as np
import SimpleITK as sitk

try:
    import torch
except ImportError:
    torch = None

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    raise SystemExit("Install openpyxl with: python -m pip install openpyxl")


PROJECT_DIR = Path(r"C:\Users\chand\BrainTumorProject\Brain-Tumors-Segmentation")
TEST_DIR = Path(r"C:\Users\chand\BrainTumorProject\dataset\brats2023\test")
OUTPUT_DIR = PROJECT_DIR / "evaluation_results"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

CSV_PATH = OUTPUT_DIR / "ground_truth_vs_prediction.csv"
XLSX_PATH = OUTPUT_DIR / "ground_truth_vs_prediction.xlsx"

if str(PROJECT_DIR) not in sys.path:
    sys.path.insert(0, str(PROJECT_DIR))

from predict import load_model, predict_patient, DEVICE


def dice_score(pred_mask, gt_mask):
    pred_mask = np.asarray(pred_mask, dtype=bool)
    gt_mask = np.asarray(gt_mask, dtype=bool)
    ps = int(pred_mask.sum())
    gs = int(gt_mask.sum())
    if ps == 0 and gs == 0:
        return 1.0
    if ps == 0 or gs == 0:
        return 0.0
    inter = int(np.logical_and(pred_mask, gt_mask).sum())
    return 2.0 * inter / (ps + gs)


def volume_cm3(voxels, spacing):
    return float(voxels) * float(spacing[0]) * float(spacing[1]) * float(spacing[2]) / 1000.0


def classification(gt_tumor, pred_tumor):
    if gt_tumor and pred_tumor:
        return "TRUE POSITIVE"
    if gt_tumor and not pred_tumor:
        return "FALSE NEGATIVE"
    if not gt_tumor and pred_tumor:
        return "FALSE POSITIVE"
    return "TRUE NEGATIVE"


def main():
    if not TEST_DIR.exists():
        raise FileNotFoundError(f"Test dataset folder not found:\n{TEST_DIR}")

    patients = sorted(p for p in TEST_DIR.iterdir() if p.is_dir())
    if not patients:
        raise RuntimeError(f"No patient folders found in {TEST_DIR}")

    print("=" * 80)
    print("GROUND TRUTH vs SEGRESNET PREDICTION")
    print("=" * 80)
    print("Patients found:", len(patients))
    print("Device:", DEVICE)
    print()

    model = load_model()
    rows = []
    tp = tn = fp = fn = 0

    for i, patient_dir in enumerate(patients, 1):
        patient_id = patient_dir.name
        seg_path = patient_dir / f"{patient_id}-seg.nii.gz"

        print(f"[{i}/{len(patients)}] {patient_id}")

        try:
            gt_img = sitk.ReadImage(str(seg_path))
            gt = sitk.GetArrayFromImage(gt_img).astype(np.uint8)

            result = predict_patient(str(patient_dir), model)
            pred = np.asarray(result["prediction"], dtype=np.uint8)

            if gt.shape != pred.shape:
                raise ValueError(f"Shape mismatch: GT={gt.shape}, Prediction={pred.shape}")

            spacing = gt_img.GetSpacing()
            size = gt_img.GetSize()

            # BraTS regions:
            # WT = 1,2,3
            # TC = 1,3
            # ET = 3
            gt_wt = gt > 0
            gt_tc = np.logical_or(gt == 1, gt == 3)
            gt_et = gt == 3

            pred_wt = pred > 0
            pred_tc = np.logical_or(pred == 1, pred == 3)
            pred_et = pred == 3

            gt_tumor_vox = int(gt_wt.sum())
            pred_tumor_vox = int(pred_wt.sum())

            gt_has = gt_tumor_vox > 0
            pred_has = pred_tumor_vox > 0
            result_class = classification(gt_has, pred_has)

            if result_class == "TRUE POSITIVE":
                tp += 1
            elif result_class == "TRUE NEGATIVE":
                tn += 1
            elif result_class == "FALSE POSITIVE":
                fp += 1
            else:
                fn += 1

            gt_edema = int((gt == 2).sum())
            gt_ncr = int((gt == 1).sum())
            gt_et_vox = int((gt == 3).sum())

            pred_edema = int((pred == 2).sum())
            pred_ncr = int((pred == 1).sum())
            pred_et_vox = int((pred == 3).sum())

            wt_dice = dice_score(pred_wt, gt_wt)
            tc_dice = dice_score(pred_tc, gt_tc)
            et_dice = dice_score(pred_et, gt_et)
            mean_dice = (wt_dice + tc_dice + et_dice) / 3.0

            rows.append({
                "Patient ID": patient_id,
                "Ground Truth": "TUMOR PRESENT" if gt_has else "NO TUMOR",
                "Prediction": "TUMOR DETECTED" if pred_has else "NO TUMOR DETECTED",
                "Detection Result": result_class,
                "GT Tumor Voxels": gt_tumor_vox,
                "Predicted Tumor Voxels": pred_tumor_vox,
                "GT Tumor Volume (cm3)": round(volume_cm3(gt_tumor_vox, spacing), 4),
                "Predicted Tumor Volume (cm3)": round(volume_cm3(pred_tumor_vox, spacing), 4),
                "GT Edema Voxels": gt_edema,
                "Predicted Edema Voxels": pred_edema,
                "GT Edema Volume (cm3)": round(volume_cm3(gt_edema, spacing), 4),
                "Predicted Edema Volume (cm3)": round(volume_cm3(pred_edema, spacing), 4),
                "GT NCR/NET Voxels": gt_ncr,
                "Predicted NCR/NET Voxels": pred_ncr,
                "GT NCR/NET Volume (cm3)": round(volume_cm3(gt_ncr, spacing), 4),
                "Predicted NCR/NET Volume (cm3)": round(volume_cm3(pred_ncr, spacing), 4),
                "GT Enhancing Tumor Voxels": gt_et_vox,
                "Predicted Enhancing Tumor Voxels": pred_et_vox,
                "GT Enhancing Tumor Volume (cm3)": round(volume_cm3(gt_et_vox, spacing), 4),
                "Predicted Enhancing Tumor Volume (cm3)": round(volume_cm3(pred_et_vox, spacing), 4),
                "WT Dice": round(wt_dice, 6),
                "TC Dice": round(tc_dice, 6),
                "ET Dice": round(et_dice, 6),
                "Overall Mean Dice": round(mean_dice, 6),
                "Voxel Spacing X (mm)": round(float(spacing[0]), 5),
                "Voxel Spacing Y (mm)": round(float(spacing[1]), 5),
                "Voxel Spacing Z (mm)": round(float(spacing[2]), 5),
                "Image Size X": int(size[0]),
                "Image Size Y": int(size[1]),
                "Image Size Z": int(size[2]),
            })

            print(f"    GT: {rows[-1]['Ground Truth']}")
            print(f"    Prediction: {rows[-1]['Prediction']}")
            print(f"    {result_class}")
            print(f"    WT={wt_dice:.4f}, TC={tc_dice:.4f}, ET={et_dice:.4f}")

        except Exception as exc:
            print(f"    ERROR: {exc}")

        finally:
            if torch is not None and torch.cuda.is_available():
                torch.cuda.empty_cache()
            gc.collect()

    if not rows:
        raise RuntimeError("No patients were successfully evaluated.")

    fields = list(rows[0].keys())

    with open(CSV_PATH, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)

    mean_wt = float(np.mean([r["WT Dice"] for r in rows]))
    mean_tc = float(np.mean([r["TC Dice"] for r in rows]))
    mean_et = float(np.mean([r["ET Dice"] for r in rows]))
    mean_all = float(np.mean([r["Overall Mean Dice"] for r in rows]))

    summary = [
        ("Patients evaluated", len(rows)),
        ("Ground Truth Tumor Present", sum(r["Ground Truth"] == "TUMOR PRESENT" for r in rows)),
        ("Ground Truth No Tumor", sum(r["Ground Truth"] == "NO TUMOR" for r in rows)),
        ("Predicted Tumor Detected", sum(r["Prediction"] == "TUMOR DETECTED" for r in rows)),
        ("Predicted No Tumor", sum(r["Prediction"] == "NO TUMOR DETECTED" for r in rows)),
        ("True Positive", tp),
        ("False Negative", fn),
        ("False Positive", fp),
        ("True Negative", tn),
        ("Mean WT Dice", mean_wt),
        ("Mean TC Dice", mean_tc),
        ("Mean ET Dice", mean_et),
        ("Overall Mean Dice", mean_all),
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = "Patient Results"
    ws.append(fields)
    for row in rows:
        ws.append([row[f] for f in fields])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    sws = wb.create_sheet("Summary")
    sws.append(["Metric", "Value"])
    for item in summary:
        sws.append(list(item))
    for cell in sws[1]:
        cell.font = Font(bold=True)

    cws = wb.create_sheet("Confusion Matrix")
    cws.append(["", "Prediction: Tumor", "Prediction: No Tumor"])
    cws.append(["GT: Tumor", tp, fn])
    cws.append(["GT: No Tumor", fp, tn])
    for cell in cws[1]:
        cell.font = Font(bold=True)

    dws = wb.create_sheet("Dice Summary")
    dws.append(["Region", "Mean Dice"])
    dws.append(["WT (Whole Tumor)", mean_wt])
    dws.append(["TC (Tumor Core)", mean_tc])
    dws.append(["ET (Enhancing Tumor)", mean_et])
    dws.append(["Overall Mean Dice", mean_all])
    for cell in dws[1]:
        cell.font = Font(bold=True)

    for sheet in wb.worksheets:
        for col in sheet.columns:
            width = min(max(len(str(c.value or "")) for c in col) + 2, 42)
            sheet.column_dimensions[get_column_letter(col[0].column)].width = max(width, 12)

    wb.save(XLSX_PATH)

    print()
    print("=" * 80)
    print("EVALUATION COMPLETE")
    print("=" * 80)
    print("Patients evaluated:", len(rows))
    print("True Positive :", tp)
    print("False Negative:", fn)
    print("False Positive:", fp)
    print("True Negative :", tn)
    print()
    print(f"Mean WT Dice: {mean_wt:.6f}")
    print(f"Mean TC Dice: {mean_tc:.6f}")
    print(f"Mean ET Dice: {mean_et:.6f}")
    print(f"Overall Mean Dice: {mean_all:.6f}")
    print()
    print("CSV :", CSV_PATH)
    print("Excel:", XLSX_PATH)


if __name__ == "__main__":
    main()
